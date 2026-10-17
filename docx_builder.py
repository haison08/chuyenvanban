import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=40, bottom=40, left=50, right=50):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, b_val in borders.items():
        if b_val:
            node = OxmlElement(f'w:{b_name}')
            for k, v in b_val.items():
                node.set(qn(f'w:{k}'), str(v))
            tcBorders.append(node)
        else:
            node = OxmlElement(f'w:{b_name}')
            node.set(qn('w:val'), 'none')
            tcBorders.append(node)
    tcPr.append(tcBorders)

def set_table_borders(table, val='single', sz='4', color='000000'):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def parse_align(align_str):
    a = str(align_str).lower().strip()
    if a == 'center':
        return WD_ALIGN_PARAGRAPH.CENTER
    elif a in ('right', 'end'):
        return WD_ALIGN_PARAGRAPH.RIGHT
    elif a in ('justify', 'justified', 'both'):
        return WD_ALIGN_PARAGRAPH.JUSTIFY
    return WD_ALIGN_PARAGRAPH.LEFT

def add_styled_paragraph(container, item, default_font='Times New Roman', default_size=13):
    if isinstance(item, str):
        item = {'text': item}
    p = container.add_paragraph()
    align = parse_align(item.get('align', 'left'))
    p.alignment = align
    
    p.paragraph_format.space_before = Pt(item.get('space_before', 0))
    p.paragraph_format.space_after = Pt(item.get('space_after', 2))
    p.paragraph_format.line_spacing = item.get('line_spacing', 1.15)
    
    first_indent = item.get('first_indent', 0)
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(first_indent if isinstance(first_indent, (int, float)) else 28)
        
    runs = item.get('runs', [])
    if not runs and 'text' in item:
        runs = [{
            'text': item['text'],
            'bold': item.get('bold', False),
            'italic': item.get('italic', False),
            'size': item.get('size', default_size),
            'font': item.get('font', default_font)
        }]
        
    for r_data in runs:
        text = r_data.get('text', '')
        if not text:
            continue
        run = p.add_run(text)
        run.bold = r_data.get('bold', False)
        run.italic = r_data.get('italic', False)
        run.font.name = r_data.get('font', default_font)
        run.font.size = Pt(r_data.get('size', default_size))
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def build_docx_from_structure(doc_data, output_path, options=None):
    if options is None:
        options = {}
        
    remove_signatures = options.get('remove_signatures', True)
    default_font = options.get('font_name', 'Times New Roman')
    
    doc = Document()
    
    # Configure base Normal style
    normal_style = doc.styles['Normal']
    normal_style.font.name = default_font
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    
    pages = doc_data.get('pages', [])
    current_orientation = 'portrait'
    
    # Configure first section
    sec = doc.sections[0]
    
    for p_idx, page in enumerate(pages):
        orientation = page.get('orientation', 'portrait').lower()
        
        # Auto-detect wide tables: if any table has >= 7 columns, force landscape
        for elem in page.get('elements', []):
            if elem.get('type') == 'table':
                r_data = elem.get('rows', [])
                n_c = len(elem.get('col_widths', [])) or (len(r_data[0].get('cells', [])) if r_data else 0)
                if n_c >= 7:
                    orientation = 'landscape'
                    break
        
        # Apply section settings
        if p_idx == 0:
            current_orientation = orientation
            if orientation == 'landscape':
                sec.orientation = WD_ORIENT.LANDSCAPE
                sec.page_width = Mm(297)
                sec.page_height = Mm(210)
                sec.top_margin = Mm(15)
                sec.bottom_margin = Mm(15)
                sec.left_margin = Mm(15)
                sec.right_margin = Mm(15)
            else:
                sec.orientation = WD_ORIENT.PORTRAIT
                sec.page_width = Mm(210)
                sec.page_height = Mm(297)
                sec.top_margin = Mm(18)
                sec.bottom_margin = Mm(18)
                sec.left_margin = Mm(22)
                sec.right_margin = Mm(15)
        elif orientation != current_orientation:
            current_orientation = orientation
            sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
            if orientation == 'landscape':
                sec.orientation = WD_ORIENT.LANDSCAPE
                sec.page_width = Mm(297)
                sec.page_height = Mm(210)
                sec.top_margin = Mm(15)
                sec.bottom_margin = Mm(15)
                sec.left_margin = Mm(15)
                sec.right_margin = Mm(15)
            else:
                sec.orientation = WD_ORIENT.PORTRAIT
                sec.page_width = Mm(210)
                sec.page_height = Mm(297)
                sec.top_margin = Mm(18)
                sec.bottom_margin = Mm(18)
                sec.left_margin = Mm(22)
                sec.right_margin = Mm(15)
                
        elements = page.get('elements', [])
        for elem in elements:
            el_type = elem.get('type')
            
            if el_type == 'header_table':
                left_items = elem.get('left', [])
                right_items = elem.get('right', [])
                tbl = doc.add_table(rows=1, cols=2)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = False
                c_l, c_r = tbl.cell(0, 0), tbl.cell(0, 1)
                
                is_landscape = (current_orientation == 'landscape')
                c_l.width = Mm(125 if is_landscape else 78)
                c_r.width = Mm(142 if is_landscape else 95)
                set_cell_margins(c_l, top=0, bottom=0, left=0, right=0)
                set_cell_margins(c_r, top=0, bottom=0, left=0, right=0)
                
                for item in left_items:
                    add_styled_paragraph(c_l, item, default_font=default_font)
                for item in right_items:
                    add_styled_paragraph(c_r, item, default_font=default_font)
                    
            elif el_type in ('paragraph', 'heading'):
                add_styled_paragraph(doc, elem, default_font=default_font)
                
            elif el_type == 'maket_box':
                box_items = elem.get('items', [])
                box_tbl = doc.add_table(rows=1, cols=1)
                box_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                c = box_tbl.cell(0, 0)
                c.width = Mm(145)
                set_cell_margins(c, top=60, bottom=60, left=120, right=120)
                set_cell_borders(c, 
                                 top={'val': 'single', 'sz': '6', 'color': '000000'},
                                 bottom={'val': 'single', 'sz': '6', 'color': '000000'},
                                 left={'val': 'single', 'sz': '6', 'color': '000000'},
                                 right={'val': 'single', 'sz': '6', 'color': '000000'})
                for item in box_items:
                    add_styled_paragraph(c, item, default_font=default_font)
                    
            elif el_type == 'signature_block':
                sig_tbl = doc.add_table(rows=1, cols=2)
                sig_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                sig_tbl.autofit = False
                c_nn, c_sig = sig_tbl.cell(0, 0), sig_tbl.cell(0, 1)
                c_nn.width = Mm(75)
                c_sig.width = Mm(95)
                set_cell_margins(c_nn, top=0, bottom=0, left=0, right=0)
                set_cell_margins(c_sig, top=0, bottom=0, left=0, right=0)
                
                # Recipients
                for item in elem.get('recipients', []):
                    add_styled_paragraph(c_nn, item, default_font=default_font)
                    
                # Signatory
                signatory_items = elem.get('signatory', [])
                for item in signatory_items:
                    if item.get('is_space'):
                        spacing = item.get('space', 35)
                        add_styled_paragraph(c_sig, {'text': '', 'space_before': spacing, 'space_after': 0}, default_font=default_font)
                    else:
                        add_styled_paragraph(c_sig, item, default_font=default_font)
                        
            elif el_type == 'table':
                rows_data = elem.get('rows', [])
                col_widths = elem.get('col_widths', [])
                num_cols = len(col_widths) if col_widths else (len(rows_data[0].get('cells', [])) if rows_data else 1)
                num_rows = len(rows_data)
                
                is_landscape = (current_orientation == 'landscape')
                available_width_mm = 267 if is_landscape else 173
                
                # Table font size based on column density
                if num_cols >= 11:
                    cell_font_size = 9.0
                elif num_cols >= 8:
                    cell_font_size = 9.5
                elif num_cols >= 6:
                    cell_font_size = 10.5
                else:
                    cell_font_size = 12.0
                    
                # Calculate intelligent column widths if not provided or mismatch
                if not col_widths or len(col_widths) != num_cols:
                    if num_cols == 12:
                        # Standard 12-column Vietnamese member/cadre list
                        col_widths = [8.5, 34.0, 18.0, 23.5, 38.0, 13.0, 13.0, 14.0, 23.5, 23.5, 22.0, 26.0]
                    else:
                        col_weights = [1.0] * num_cols
                        for r in rows_data:
                            for c_i, cell in enumerate(r.get('cells', [])):
                                c_idx = cell.get('col', c_i)
                                txt_len = len(str(cell.get('text', '')).strip())
                                if c_idx < num_cols:
                                    col_weights[c_idx] = max(col_weights[c_idx], txt_len)
                        
                        # Column 0 is usually index/TT -> cap weight
                        if num_cols >= 6:
                            col_weights[0] = min(col_weights[0], 4)
                        
                        total_weight = sum(col_weights) or 1
                        col_widths = [max(8, round((w / total_weight) * available_width_mm, 1)) for w in col_weights]
                        scale = available_width_mm / sum(col_widths)
                        col_widths = [round(w * scale, 1) for w in col_widths]
                
                table = doc.add_table(rows=num_rows, cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False
                set_table_borders(table, val='single', sz='4', color='000000')
                
                # Apply column widths and cell padding
                for r_idx, r_elem in enumerate(table.rows):
                    for c_idx, w_mm in enumerate(col_widths):
                        if c_idx < len(r_elem.cells):
                            r_elem.cells[c_idx].width = Mm(w_mm)
                            r_elem.cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                            set_cell_margins(r_elem.cells[c_idx], top=15, bottom=15, left=20, right=20)
                            
                # Fill cells and handle merges
                for r_idx, row_obj in enumerate(rows_data):
                    cells = row_obj.get('cells', [])
                    for c_pos, cell_info in enumerate(cells):
                        c_start = cell_info.get('col', c_pos)
                        r_start = cell_info.get('row', r_idx)
                        colspan = cell_info.get('colspan', 1)
                        rowspan = cell_info.get('rowspan', 1)
                        
                        if r_start >= num_rows or c_start >= num_cols:
                            continue
                            
                        target_cell = table.cell(r_start, c_start)
                        
                        # Handle horizontal merge
                        if colspan > 1:
                            for c_offset in range(1, colspan):
                                if c_start + c_offset < num_cols:
                                    target_cell.merge(table.cell(r_start, c_start + c_offset))
                                    
                        # Handle vertical merge
                        if rowspan > 1:
                            for r_offset in range(1, rowspan):
                                if r_start + r_offset < num_rows:
                                    target_cell.merge(table.cell(r_start + r_offset, c_start))
                                    
                        # Format text with appropriate table font size
                        cell_format = dict(cell_info)
                        if 'size' not in cell_format:
                            cell_format['size'] = cell_font_size
                        if 'space_before' not in cell_format:
                            cell_format['space_before'] = 0
                        if 'space_after' not in cell_format:
                            cell_format['space_after'] = 1
                        if 'line_spacing' not in cell_format:
                            cell_format['line_spacing'] = 1.05
                            
                        text = cell_format.get('text', '')
                        if text:
                            add_styled_paragraph(target_cell, cell_format, default_font=default_font)
                            
                    # Custom row height if specified
                    if 'height_val' in row_obj:
                        h_val = row_obj['height_val']
                        trPr = table.rows[r_idx]._tr.get_or_add_trPr()
                        trHeight = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="{h_val}" w:hRule="atLeast"/>')
                        trPr.append(trHeight)
                        
    # End paragraph to prevent trailing empty page
    end_p = doc.add_paragraph()
    end_p.paragraph_format.space_before = Pt(0)
    end_p.paragraph_format.space_after = Pt(0)
    end_p.paragraph_format.line_spacing = Pt(1)
    end_run = end_p.add_run("")
    end_run.font.size = Pt(1)
    
    doc.save(output_path)
    return output_path
