import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=40, bottom=40, left=50, right=50):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, b_val in borders.items():
        if b_val:
            node = OxmlElement(f'w:{b_name}')
            for k, v in b_val.items():
                node.set(qn(f'w:{k}'), str(v))
            tcBorders.append(node)
        else:
            node = OxmlElement(f'w:{b_name}')
            node.set(qn('w:val'), 'none')
            tcBorders.append(node)
    tcPr.append(tcBorders)

def set_table_borders(table, val='single', sz='4', color='000000'):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def add_p(container, text="", bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, 
          size=13, space_before=0, space_after=1.5, line_spacing=1.12, first_indent=0):
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if first_indent > 0:
        p.paragraph_format.first_line_indent = Pt(first_indent)
    
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def generate_document():
    doc = Document()
    
    # Base Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    
    # Section 1: Portrait (Pages 1-5)
    # Calibrated margins so pages 1 to 5 flow naturally with zero spillover
    section1 = doc.sections[0]
    section1.page_width = Mm(210)
    section1.page_height = Mm(297)
    section1.top_margin = Mm(18)
    section1.bottom_margin = Mm(18)
    section1.left_margin = Mm(22)
    section1.right_margin = Mm(15)
    
    # =========================================================================
    # HEADER (Quốc hiệu & Đơn vị ban hành)
    # =========================================================================
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    
    cell_left = header_table.cell(0, 0)
    cell_right = header_table.cell(0, 1)
    cell_left.width = Mm(78)
    cell_right.width = Mm(95)
    set_cell_margins(cell_left, top=0, bottom=0, left=0, right=0)
    set_cell_margins(cell_right, top=0, bottom=0, left=0, right=0)
    
    # Left Header
    add_p(cell_left, "ỦY BAN MTTQ VIỆT NAM", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=11.5, space_after=1, line_spacing=1.05)
    add_p(cell_left, "PHƯỜNG THÀNH SEN", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=11.5, space_after=1, line_spacing=1.05)
    add_p(cell_left, "BAN THƯỜNG TRỰC", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=1, line_spacing=1.05)
    add_p(cell_left, "——————", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=9, space_after=1, line_spacing=1.0)
    add_p(cell_left, "Số: 20/KH-MTTQ-BTT", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=0, line_spacing=1.05)
    
    # Right Header
    add_p(cell_right, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11.5, space_after=1, line_spacing=1.05)
    add_p(cell_right, "Độc lập - Tự do - Hạnh phúc", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12.5, space_after=1, line_spacing=1.05)
    add_p(cell_right, "————————————", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=9, space_after=2, line_spacing=1.0)
    add_p(cell_right, "Thành Sen, ngày 05 tháng 10 năm 2026", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12.5, space_after=0, line_spacing=1.05)
    
    # Tiêu đề văn bản
    add_p(doc, "KẾ HOẠCH", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_before=4, space_after=2)
    add_p(doc, "Tổ chức Ngày hội Đại đoàn kết toàn dân tộc ở khu dân cư năm 2026", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13.5, space_before=0, space_after=6)
    
    # Căn cứ
    p_cc = doc.add_paragraph()
    p_cc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_cc.paragraph_format.first_line_indent = Pt(28)
    p_cc.paragraph_format.line_spacing = 1.15
    p_cc.paragraph_format.space_before = Pt(0)
    p_cc.paragraph_format.space_after = Pt(2)
    
    r1 = p_cc.add_run("Thực hiện Hướng dẫn số 18/HD-MTTQ-BTT ngày 24/9/2026 của Ủy ban MTTQ Việt Nam tỉnh Hà Tĩnh về việc tổ chức Ngày hội Đại đoàn kết toàn dân tộc ở khu dân cư giai đoạn 2026-2031; và thực hiện các chủ trương Nghị quyết Đại hội đại biểu Đảng bộ các cấp nhiệm kỳ 2025-2030; Nghị quyết Đại hội đại biểu MTTQ Việt Nam các cấp, nhiệm kỳ 2025 - 2030 về phát huy sức mạnh đại đoàn kết toàn dân tộc, xây dựng đời sống hóa ở khu dân cư, xây dựng ")
    r1.font.size = Pt(13)
    
    r_quote = p_cc.add_run("“nông thôn mới, đô thị văn minh”")
    r_quote.italic = True
    r_quote.font.size = Pt(13)
    
    r2 = p_cc.add_run(" và nâng cao đời sống Nhân dân; Kỷ niệm 96 năm Ngày truyền thống MTTQ Việt Nam (18/11/1930 - 18/11/2026); Ban Thường trực Ủy ban MTTQVN phường Thành Sen xây dựng kế hoạch triển khai thực hiện như sau:")
    r2.font.size = Pt(13)
    
    # I. MỤC ĐÍCH, YÊU CẦU
    add_p(doc, "I. MỤC ĐÍCH, YÊU CẦU", bold=True, size=13, space_before=3, space_after=2, first_indent=28)
    add_p(doc, "1. Mục đích", bold=True, size=13, space_before=1, space_after=2, first_indent=28)
    
    add_p(doc, "- Cụ thể hóa chủ trương của Đảng, chính sách, pháp luật của Nhà nước về phát huy sức mạnh đại đoàn kết toàn dân tộc, xây dựng đời sống văn hóa và phát huy vai trò chủ thể của Nhân dân; triển khai thực hiện hiệu quả Nghị quyết Đại hội đại biểu Đảng bộ các cấp nhiệm kỳ 2025 - 2030 và Nghị quyết Đại hội đại biểu MTTQ Việt Nam các cấp, nhiệm kỳ 2025 - 2030.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "- Tiếp tục phát huy truyền thống vẻ vang của Mặt trận Dân tộc thống nhất Việt Nam; củng cố, tăng cường khối đại đoàn kết toàn dân, tạo sự đồng thuận xã hội, phát huy quyền làm chủ, tinh thần tự quản, ý thức trách nhiệm của Nhân dân trong xây dựng quê hương Hà Tĩnh phát triển nhanh, bền vững.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "- Thông qua Ngày hội nhằm tuyên truyền, vận động Nhân dân thực hiện các chủ trương, đường lối của Đảng, chính sách pháp luật của Nhà nước; đánh giá kết quả thực hiện các cuộc vận động, phong trào thi đua yêu nước; biểu dương, tôn vinh các tập thể, gia đình, cá nhân tiêu biểu, những cách làm hay, mô hình hiệu quả, những người có nhiều đóng góp cho cộng đồng.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "2. Yêu cầu", bold=True, size=13, space_before=1, space_after=2, first_indent=28)
    
    add_p(doc, "- Ban Thường trực Uỷ ban MTTQ phường chỉ đạo Ban Công tác Mặt trận phối hợp các đồng chí trong liên đoàn cán bộ báo cáo cấp ủy phối hợp tổ chức thực hiện. Ngày hội “Đại đoàn kết toàn dân tộc” được tổ chức ở tất cả cộng đồng, khu dân cư trong toàn phường, huy động sự tham gia đông đảo các tầng lớp nhân dân tham gia. Việc tổ chức Ngày hội đảm bảo thiết thực tiết kiệm, an toàn, tránh hình thức, lãng phí.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "- Nội dung phần Lễ tổ chức trang trọng, ngắn gọn; nội dung phần Hội tạo được không khí vui tươi, đầm ấm và tôn vinh giá trị văn hóa truyền thống của dân tộc, tín ngưỡng tôn giáo, vùng, miền.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    # II. NỘI DUNG CỤ THỂ
    add_p(doc, "II. NỘI DUNG CỤ THỂ", bold=True, size=13, space_before=3, space_after=2, first_indent=28)
    add_p(doc, "1. Công tác tuyên truyền", bold=True, size=13, space_before=1, space_after=2, first_indent=28)
    add_p(doc, "1.1. Nội dung tuyên truyền", bold=True, italic=True, size=13, space_before=1, space_after=2, first_indent=28)
    
    add_p(doc, "- Truyền thống lịch sử vẻ vang của Mặt trận Tổ quốc Việt Nam qua các thời kỳ; ý nghĩa, vai trò của Ngày hội “Đại đoàn kết toàn dân tộc”.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "- Tuyên truyền chủ trương, đường lối của Đảng, chính sách, pháp luật của Nhà nước; Nghị quyết Đại hội đại biểu toàn quốc lần thứ XIV của Đảng; Nghị quyết Đại hội đại biểu Đảng bộ các cấp, nhiệm kỳ 2025 - 2030; Nghị quyết Đại hội đại biểu MTTQ Việt Nam các cấp nhiệm kỳ 2025 - 2030. Ôn lại truyền thống 96 năm hình thành và phát triển của Mặt trận Tổ quốc Việt Nam (18/11/1930 - 18/11/2026) và kết quả xây dựng khối đại đoàn kết toàn dân tộc.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "- Tuyên truyền kết quả thực hiện nội dung các cuộc vận động, phong trào thi đua yêu nước do MTTQ Việt Nam các cấp chủ trì, phối hợp thực hiện.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "- Tuyên truyền gương người tốt, việc tốt; các tập thể, cá nhân tiêu biểu trong thực hiện các cuộc vận động, các phong trào thi đua và nhiệm vụ chính trị của địa phương.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "1.2. Hình thức, nội dung tuyên truyền", bold=True, italic=True, size=13, space_before=1, space_after=2, first_indent=28)
    
    add_p(doc, "- Tuyên truyền trên hệ thống phát thanh, truyền hình, báo chí, hệ thống truyền thanh cơ sở, cổng thông tin điện tử, trang Fanpage, mạng xã hội của MTTQ Việt Nam các cấp, các tổ chức thành viên, các cơ quan, đơn vị, tổ chức và mỗi cá nhân.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "- Tuyên truyền thông qua cuộc họp, sinh hoạt của Nhân dân, sinh hoạt của các tổ chức đoàn thể tại khu dân cư.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "- Tổ chức treo băng rôn, khẩu hiệu, pano, áp phích, cờ, hoa tại khu dân cư, địa điểm tổ chức Ngày hội, trung tâm văn hóa - xã hội và các cơ quan, đơn vị, trường học; tổ chức hoạt động văn hóa, văn nghệ, dân vũ, thể thao;", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "- Vận động các hộ gia đình treo cờ Tổ quốc trong dịp tổ chức Ngày hội từ ngày 01/11/2026 đến hết ngày 18/11/2026.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    # 2. Khẩu hiệu tuyên truyền
    add_p(doc, "2. Khẩu hiệu tuyên truyền", bold=True, size=13, space_before=1, space_after=2, first_indent=28)
    
    slogans = [
        "- Kỷ niệm 96 năm Ngày truyền thống MTTQ Việt Nam (18/11/1930 - 18/11/2026)!",
        "- Phát huy sức mạnh đại đoàn kết toàn dân tộc vì sự nghiệp dân giàu, nước mạnh, dân chủ, công bằng, văn minh!",
        "- Nhiệt liệt chào mừng Ngày hội đại đoàn kết toàn dân tộc!",
        "- Sống, chiến đấu, lao động, học tập và làm theo tư tưởng, đạo đức, phong cách Hồ Chí Minh!",
        "- Chủ tịch Hồ Chí Minh vĩ đại sống mãi trong sự nghiệp của chúng ta!",
        "- Đoàn kết, đoàn kết, đại đoàn kết. Thành công, thành công, đại thành công!",
        "- Toàn dân đoàn kết xây dựng nông thôn mới đô thị văn minh.",
        "- Quyết tâm xây dựng phường Thành Sen ngày càng giàu đẹp văn minh."
    ]
    for s in slogans:
        add_p(doc, s, italic=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
        
    # 3. Tổ chức Ngày hội Đại đoàn kết
    add_p(doc, "3. Tổ chức Ngày hội Đại đoàn kết", bold=True, size=13, space_before=2, space_after=2, first_indent=28)
    add_p(doc, "3.1. Thời gian", bold=True, italic=True, size=13, space_before=1, space_after=2, first_indent=28)
    
    add_p(doc, "- Ngày hội và các hoạt động tổ chức Ngày hội diễn ra cao điểm từ ngày 01/11 đến hết ngày 18/11/2026.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    p_tg2 = doc.add_paragraph()
    p_tg2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_tg2.paragraph_format.first_line_indent = Pt(28)
    p_tg2.paragraph_format.line_spacing = 1.12
    p_tg2.paragraph_format.space_before = Pt(0)
    p_tg2.paragraph_format.space_after = Pt(2)
    r = p_tg2.add_run("- Các tổ dân phố đăng ký lịch tổ chức Ngày hội về Ban Thường trực Uỷ ban MTTQ phường trước ngày 20/10/2026 ")
    r.font.size = Pt(13)
    r_it = p_tg2.add_run("(Qua đồng chí Nguyễn Thị Mai, SĐT: 0912.488235)")
    r_it.italic = True
    r_it.font.size = Pt(13)
    r_end = p_tg2.add_run(". Trên cơ sở lịch đã đăng ký, Ban Thường trực Uỷ ban MTTQ phường sắp xếp thời gian đảm bảo phù hợp, tránh chồng chéo để các đoàn công tác của Ban Thường vụ Đảng uỷ tham dự.")
    r_end.font.size = Pt(13)
    
    add_p(doc, "3.2. Hình thức và chủ trì tổ chức Ngày hội", bold=True, italic=True, size=13, space_before=1, space_after=2, first_indent=28)
    add_p(doc, "- Hình thức tổ chức: Ngày hội được tổ chức ở địa bàn khu dân cư; tùy theo điều kiện ở mỗi khu dân cư có thể lựa chọn hình thức phù hợp.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    add_p(doc, "- Chủ trì tổ chức Ngày hội: Do Ban công tác Mặt trận chủ trì, phối hợp với tổ dân phố.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "3.3. Thành phần tham gia Ngày hội", bold=True, italic=True, size=13, space_before=1, space_after=2, first_indent=28)
    add_p(doc, "- Liên đoàn cán bộ Tổ dân phố và toàn thể Nhân dân hoặc đại diện hộ gia đình ở khu dân cư.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    add_p(doc, "- Mời đại diện lãnh đạo các cơ quan, đơn vị, lực lượng vũ trang đóng trên địa bàn tổ dân phố; đại diện các tổ chức, doanh nghiệp, cá nhân có hoạt động hỗ trợ, đóng góp cho cộng đồng dân cư.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    add_p(doc, "- Mời đại diện Đoàn công tác của Ban Thường vụ Đảng ủy chỉ đạo các TDP; đại diện lãnh đạo cấp ủy, chính quyền, Mặt trận Tổ quốc và các đoàn thể chính trị - xã hội, cán bộ công chức phường.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "3.4. Hình thức trang trí", bold=True, italic=True, size=13, space_before=1, space_after=2, first_indent=28)
    add_p(doc, "- Treo cờ, băng rôn, khẩu hiệu tại Nhà văn hóa, các điểm công cộng, các tuyến đường chính trên địa bàn khu dân cư và địa điểm tổ chức các hoạt động Ngày hội.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    p_mk = doc.add_paragraph()
    p_mk.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_mk.paragraph_format.first_line_indent = Pt(28)
    p_mk.paragraph_format.line_spacing = 1.12
    p_mk.paragraph_format.space_before = Pt(0)
    p_mk.paragraph_format.space_after = Pt(3)
    r_mk1 = p_mk.add_run("- Ma két trang trí nơi tổ chức phần Lễ của Ngày hội: Đối với địa bàn dân cư ")
    r_mk1.font.size = Pt(13)
    r_mk2 = p_mk.add_run("(do Ban Công tác Mặt trận chủ trì)")
    r_mk2.italic = True
    r_mk2.font.size = Pt(13)
    r_mk3 = p_mk.add_run(":")
    r_mk3.font.size = Pt(13)
    
    # KHUNG MA KET (Single-cell Table)
    maket_tbl = doc.add_table(rows=1, cols=1)
    maket_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_mk = maket_tbl.cell(0, 0)
    cell_mk.width = Mm(145)
    set_cell_margins(cell_mk, top=60, bottom=60, left=120, right=120)
    set_cell_borders(cell_mk, 
                     top={'val': 'single', 'sz': '6', 'color': '000000'},
                     bottom={'val': 'single', 'sz': '6', 'color': '000000'},
                     left={'val': 'single', 'sz': '6', 'color': '000000'},
                     right={'val': 'single', 'sz': '6', 'color': '000000'})
    
    add_p(cell_mk, "Cờ Tổ quốc, ảnh Bác", align=WD_ALIGN_PARAGRAPH.LEFT, size=12, space_after=1, line_spacing=1.05)
    add_p(cell_mk, "(Logo Mặt trận Tổ quốc Việt Nam)", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=6, line_spacing=1.05)
    add_p(cell_mk, "NGÀY HỘI", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13, space_after=1, line_spacing=1.05)
    add_p(cell_mk, "ĐẠI ĐOÀN KẾT TOÀN DÂN TỘC", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13.5, space_after=1, line_spacing=1.05)
    add_p(cell_mk, "Tổ dân phố..., phường Thành Sen", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12.5, space_after=8, line_spacing=1.05)
    add_p(cell_mk, "Thành Sen, ngày ... tháng 11 năm 2026", italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT, size=12, space_after=1, line_spacing=1.05)
    
    add_p(doc, "(Lưu ý: maket nền đỏ chữ vàng, có thể phối hình hoa sen phía dưới)", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_before=2, space_after=3)
    
    # 3.5. Chương trình Ngày hội
    add_p(doc, "3.5. Chương trình Ngày hội", bold=True, italic=True, size=13, space_before=2, space_after=2, first_indent=28)
    
    p_le = doc.add_paragraph()
    p_le.paragraph_format.first_line_indent = Pt(28)
    p_le.paragraph_format.space_before = Pt(0)
    p_le.paragraph_format.space_after = Pt(2)
    p_le.paragraph_format.line_spacing = 1.12
    r_le1 = p_le.add_run("* Phần Lễ ")
    r_le1.bold = True
    r_le1.font.size = Pt(13)
    r_le2 = p_le.add_run("(Tổ chức không quá 90 phút)")
    r_le2.italic = True
    r_le2.font.size = Pt(13)
    r_le3 = p_le.add_run(":")
    r_le3.bold = True
    r_le3.font.size = Pt(13)
    
    p_vn = doc.add_paragraph()
    p_vn.paragraph_format.first_line_indent = Pt(28)
    p_vn.paragraph_format.space_before = Pt(0)
    p_vn.paragraph_format.space_after = Pt(2)
    p_vn.paragraph_format.line_spacing = 1.12
    r_vn1 = p_vn.add_run("(1) Văn nghệ chào mừng ")
    r_vn1.font.size = Pt(13)
    r_vn2 = p_vn.add_run("(từ 4-5 tiết mục có lựa chọn chuẩn bị chu đáo)")
    r_vn2.italic = True
    r_vn2.font.size = Pt(13)
    r_vn3 = p_vn.add_run(".")
    r_vn3.font.size = Pt(13)
    
    add_p(doc, "(2) Chào cờ.", size=13, space_after=2, first_indent=28)
    add_p(doc, "(3) Tuyên bố lý do giới thiệu đại biểu", size=13, space_after=2, first_indent=28)
    add_p(doc, "(4) Ôn lại truyền thống 96 năm Ngày Thành lập Mặt trận Dân tộc thống nhất Việt Nam - Ngày truyền thống Mặt trận Tổ quốc Việt Nam; báo cáo khái quát kết quả thực hiện nhiệm vụ xây dựng khối đại đoàn kết toàn dân tộc;", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    p_bc = doc.add_paragraph()
    p_bc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_bc.paragraph_format.first_line_indent = Pt(28)
    p_bc.paragraph_format.line_spacing = 1.12
    p_bc.paragraph_format.space_before = Pt(0)
    p_bc.paragraph_format.space_after = Pt(2)
    r_bc1 = p_bc.add_run("- Kết quả thực hiện các cuộc vận động, phong trào thi đua yêu nước sau 01 năm thực hiện, trọng tâm là Cuộc vận động “Toàn dân đoàn kết xây dựng nông thôn mới, đô thị văn minh” (có nêu các tập thể, hộ gia đình, cá nhân tiêu biểu có nhiều đóng góp xây dựng cộng đồng dân cư); mục tiêu và cách thức thực hiện trong thời gian tới. ")
    r_bc1.font.size = Pt(13)
    r_bc2 = p_bc.add_run("(Khuyến khích xây dựng báo cáo bằng video, hình ảnh và ứng dụng công nghệ chuyển đổi số)")
    r_bc2.italic = True
    r_bc2.bold = True
    r_bc2.font.size = Pt(13)
    
    add_p(doc, "(5) Trao đổi, thảo luận.", size=13, space_after=2, first_indent=28)
    add_p(doc, "(6) Thông qua quyết định công nhận gia đình văn hóa năm 2026; biểu dương, tôn vinh, khen thưởng những tập thể, gia đình và cá nhân tiêu biểu.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    add_p(doc, "(7) Mời lãnh đạo cấp trên phát biểu, tặng quà (nếu có).", size=13, space_after=2, first_indent=28)
    add_p(doc, "(8) Phát động hưởng ứng phong trào thi đua xây dựng khu dân cư đoàn kết, an toàn, ấm no, hạnh phúc.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    add_p(doc, "(9) Phát biểu cảm ơn và kết thúc buổi Lễ.", size=13, space_after=2, first_indent=28)
    
    p_hoi = doc.add_paragraph()
    p_hoi.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_hoi.paragraph_format.first_line_indent = Pt(28)
    p_hoi.paragraph_format.line_spacing = 1.12
    p_hoi.paragraph_format.space_before = Pt(0)
    p_hoi.paragraph_format.space_after = Pt(2)
    r_h1 = p_hoi.add_run("* Phần Hội ")
    r_h1.bold = True
    r_h1.font.size = Pt(13)
    r_h2 = p_hoi.add_run("(Có thể tổ chức trước hoặc sau phần Lễ)")
    r_h2.italic = True
    r_h2.font.size = Pt(13)
    r_h3 = p_hoi.add_run(": Tùy điều kiện từng khu dân cư, để lựa chọn, tổ chức các hoạt động phù hợp.")
    r_h3.font.size = Pt(13)
    
    add_p(doc, "- Tổ chức các hoạt động văn hóa, thể dục, thể thao, các trò chơi dân gian truyền thống của địa phương, sinh hoạt ngoài trời, các hoạt động dân vũ...", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    add_p(doc, "- Tùy từng điều kiện các khu dân cư có thể tổ chức “Bữa cơm đoàn kết” phù hợp và ý nghĩa thiết thực.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    add_p(doc, "- Khuyến khích các khu dân cư có sáng kiến, đổi mới tổ chức phần Hội đảm bảo vui tươi, an toàn, lành mạnh, tiết kiệm, phù hợp với điều kiện, phong tục, tập quán của địa phương.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    # III. TỔ CHỨC THỰC HIỆN
    add_p(doc, "III. TỔ CHỨC THỰC HIỆN", bold=True, size=13, space_before=3, space_after=2, first_indent=28)
    add_p(doc, "1. Ban Thường trực Uỷ ban MTTQ phường", bold=True, size=13, space_before=1, space_after=2, first_indent=28)
    
    add_p(doc, "- Hướng dẫn Ban Công tác Mặt trận các khu dân cư triển khai tổ chức Ngày hội “Đại đoàn kết toàn dân tộc” phù hợp với tình hình thực tế của địa phương, đơn vị đảm bảo thời gian quy định.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    add_p(doc, "- Tham mưu Thường trực Đảng ủy phân công các đồng chí trong Ban Thường vụ, Đoàn công tác của Ban Thường vụ Đảng ủy thực hiện nhiệm vụ chỉ đạo, kiểm tra và tham dự Ngày hội tại các khu dân cư.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    add_p(doc, "- Phối hợp với Phòng Văn hoá - Xã hội tổ chức bình xét, công nhận danh hiệu “Gia đình văn hóa”; danh hiệu, tổ dân phố văn hóa để công bố và biểu dương trong dịp tổ chức Ngày hội.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "2. Ban Công tác Mặt trận các tổ dân phố", bold=True, size=13, space_before=1, space_after=2, first_indent=28)
    add_p(doc, "- Ban Công tác Mặt trận phối hợp chặt chẽ với tổ dân phố, các chi hội, chi đoàn trong tổ chức Ngày hội.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    add_p(doc, "- Tổ chức tuyên truyền, vận động Nhân dân hoàn thành chương trình mục tiêu năm 2026 theo Nghị quyết của HĐND phường; tổ chức các hoạt động hướng về kỷ niệm 96 năm ngày thành lập MTTQ Việt Nam.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    add_p(doc, "- Tổ chức tốt Ngày hội “Đại đoàn kết toàn dân tộc” và triển khai vận động ủng hộ Quỹ \"Vì người nghèo\" năm 2026 từ ngày 17/10/2026 đến 18/11/2026.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=2, first_indent=28)
    
    add_p(doc, "3. Về thông tin báo cáo", bold=True, size=13, space_before=1, space_after=2, first_indent=28)
    
    p_bc1 = doc.add_paragraph()
    p_bc1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_bc1.paragraph_format.first_line_indent = Pt(28)
    p_bc1.paragraph_format.line_spacing = 1.12
    p_bc1.paragraph_format.space_before = Pt(0)
    p_bc1.paragraph_format.space_after = Pt(2)
    r1 = p_bc1.add_run("- Ban Công tác Mặt trận các tổ dân phố: Cập nhật thông tin về tình hình tổ chức Ngày hội trước 15h00 hàng ngày (Từ ngày 05/11/2026 - 18/11/2026) và báo cáo kết quả tổ chức thực hiện ")
    r1.font.size = Pt(13)
    r_pl = p_bc1.add_run("(Có phụ lục gửi kèm)")
    r_pl.italic = True
    r_pl.font.size = Pt(13)
    r2 = p_bc1.add_run(" về Ban Thường trực Uỷ ban MTTQ phường trước ngày 19/11/2026.")
    r2.font.size = Pt(13)
    
    add_p(doc, "- Ban Thường trực Uỷ ban MTTQ phường thường xuyên cập nhật thông tin về tình hình tổ chức Ngày hội; tổng hợp tình hình, kết quả tổ chức Ngày hội báo cáo Ban Thường trực Ủy ban MTTQ Tỉnh và Ban Thường vụ Đảng ủy trước ngày 19/11/2026.", 
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_after=3, first_indent=28)
    
    add_p(doc, "Trên đây là kế hoạch tổ chức Ngày hội đại đoàn kết toàn dân tộc và các hoạt động kỷ niệm 96 năm ngày truyền thống MTTQ Việt Nam của Ủy ban MTTQ phường Thành Sen./.", 
          italic=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=13, space_before=2, space_after=8, first_indent=28)
    
    # PHẦN CHỮ KÝ & NƠI NHẬN
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    
    c_nn = sig_table.cell(0, 0)
    c_sig = sig_table.cell(0, 1)
    c_nn.width = Mm(75)
    c_sig.width = Mm(95)
    set_cell_margins(c_nn, top=0, bottom=0, left=0, right=0)
    set_cell_margins(c_sig, top=0, bottom=0, left=0, right=0)
    
    # Nơi nhận
    add_p(c_nn, "Nơi nhận:", bold=True, italic=True, size=12, space_after=1, line_spacing=1.0)
    noi_nhan = [
        "- BTT UBMTTQ tỉnh (để báo cáo);",
        "- Ban Thường vụ, Ban Chấp hành Đảng ủy;",
        "- Ban XD Đảng, UBKT, VP Đảng ủy;",
        "- Ban Thường trực UBMTTQ phường;",
        "- Ban CTMT, Tổ trưởng các TDP;",
        "- Lưu: VT."
    ]
    for nn in noi_nhan:
        add_p(c_nn, nn, size=11, space_after=1, line_spacing=1.0)
        
    # Chữ ký (khoảng cách chuẩn văn thư để ký tay và đóng dấu)
    add_p(c_sig, "TM. BAN THƯỜNG TRỰC", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13, space_after=1, line_spacing=1.05)
    add_p(c_sig, "CHỦ TỊCH", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13, space_after=1, line_spacing=1.05)
    add_p(c_sig, "", space_before=35, space_after=0)
    add_p(c_sig, "Nguyễn Thị Hồng Minh", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13.5, space_before=0, space_after=0, line_spacing=1.05)
    
    # =========================================================================
    # SECTION 2: BẢNG PHỤ LỤC (KHỔ GIẤY NGANG - A4 LANDSCAPE)
    # =========================================================================
    sec2 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec2.page_width = Mm(297)
    sec2.page_height = Mm(210)
    sec2.top_margin = Mm(14)
    sec2.bottom_margin = Mm(12)
    sec2.left_margin = Mm(14)
    sec2.right_margin = Mm(14)
    
    # Header Phụ lục
    pl_header = doc.add_table(rows=1, cols=2)
    pl_header.alignment = WD_TABLE_ALIGNMENT.CENTER
    pl_header.autofit = False
    
    c_pl_l = pl_header.cell(0, 0)
    c_pl_r = pl_header.cell(0, 1)
    c_pl_l.width = Mm(115)
    c_pl_r.width = Mm(154)
    set_cell_margins(c_pl_l, top=0, bottom=0, left=0, right=0)
    set_cell_margins(c_pl_r, top=0, bottom=0, left=0, right=0)
    
    add_p(c_pl_l, "ỦY BAN MTTQ VIỆT NAM", align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, space_after=1, line_spacing=1.0)
    add_p(c_pl_l, "PHƯỜNG THÀNH SEN", align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, space_after=1, line_spacing=1.0)
    add_p(c_pl_l, "BAN CTMT TDP ........", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11, space_after=1, line_spacing=1.0)
    
    add_p(c_pl_r, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11, space_after=1, line_spacing=1.0)
    add_p(c_pl_r, "Độc lập – Tự do – Hạnh phúc", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11.5, space_after=1, line_spacing=1.0)
    add_p(c_pl_r, "————————————", align=WD_ALIGN_PARAGRAPH.CENTER, size=8, space_after=1, line_spacing=1.0)
    
    # Tiêu đề bảng
    add_p(doc, "BẢNG TỔNG HỢP", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12.5, space_before=2, space_after=1)
    add_p(doc, "KẾT QUẢ NGÀY HỘI ĐẠI ĐOÀN KẾT TOÀN DÂN TỘC NĂM 2026", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12.5, space_after=1)
    add_p(doc, "(Kèm theo kế hoạch số 20/KH-MTTQ-BTT, ngày 05/10/2026 của Ban Thường trực Ủy ban MTTQ phường Thành Sen)", 
          italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, space_after=4)
    
    # BẢNG TỔNG HỢP 20 CỘT
    table = doc.add_table(rows=6, cols=20)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, val='single', sz='4', color='000000')
    
    col_widths = [
        16, # 0: Tổ dân phố
        13, # 1: Cả lễ & hội
        11, # 2: Lễ
        11, # 3: Hội
        13, # 4: Bữa cơm
        15, # 5: Tập thể
        15, # 6: Hộ gia đình
        13, # 7: Cá nhân
        16, # 8: Tổng KP
        15, # 9: Thăm hỏi
        13, # 10: Người tham dự
        10, # 11: Tổng số LĐ
        13, # 12: LĐ TƯ
        14, # 13: LĐ Tỉnh
        14, # 14: LĐ Sở
        14, # 15: LĐ Xã
        13, # 16: Công trình
        12, # 17: Mô hình
        15, # 18: Nhà ĐĐK mới
        15  # 19: Nhà ĐĐK sửa
    ]
    
    for row in table.rows:
        for idx, w in enumerate(col_widths):
            row.cells[idx].width = Mm(w)
            row.cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(row.cells[idx], top=25, bottom=25, left=25, right=25)
            
    # ROW 0:
    cell_h0 = table.cell(0, 1)
    for c in range(2, 20):
        cell_h0.merge(table.cell(0, c))
    add_p(cell_h0, "HOẠT ĐỘNG TỔ CHỨC NGÀY HỘI", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, space_after=0, line_spacing=1.0)
    
    cell_tdp = table.cell(0, 0)
    cell_tdp.merge(table.cell(1, 0)).merge(table.cell(2, 0))
    add_p(cell_tdp, "Tổ\ndân\nphố", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=8, space_after=0, line_spacing=1.0)
    
    # Headers in Row 1 (merged with Row 2 where no sub-headers):
    c = table.cell(1, 1).merge(table.cell(2, 1))
    add_p(c, "Tổ\nchức\ncả\nphần\nLễ và\nphần\nHội\n(Đánh\ndấu X)", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    c = table.cell(1, 2).merge(table.cell(2, 2))
    add_p(c, "Tổ\nchức\nphần\nLễ\n\n(Đánh\ndấu\nX)", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    c = table.cell(1, 3).merge(table.cell(2, 3))
    add_p(c, "Tổ\nchức\nphần\nHội\n\n(Đánh\ndấu X)", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    c = table.cell(1, 4).merge(table.cell(2, 4))
    add_p(c, "Số tổ\nliên\ngia tổ\nchức\nbữa\ncơm\nĐại\nđoàn\nkết", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    c = table.cell(1, 5).merge(table.cell(2, 5))
    add_p(c, "Số tập\nthể tiêu\nbiểu\nđược biểu\ndương,\nkhen\nthưởng\ntrong\nNgày hội/\nKinh phí", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    c = table.cell(1, 6).merge(table.cell(2, 6))
    add_p(c, "Số hộ gia\nđình tiêu\nbiểu\nđược biểu\ndương\nkhen\nthưởng\ntrong\nNgày hội/\nKinh phí", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    c = table.cell(1, 7).merge(table.cell(2, 7))
    add_p(c, "Số cá\nnhân\nđược\nbiểu\ndương,\nkhen\nthưởng\ntrong\nNgày\nhội/\nKinh phí", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    c = table.cell(1, 8).merge(table.cell(2, 8))
    add_p(c, "Tổng kinh\nphí biểu\ndương,\nkhen\nthưởng\ncho các\ntập thể, hộ\ngia đình,\ncá nhân\ntiêu biểu...\n(=6+7+8)", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    c = table.cell(1, 9).merge(table.cell(2, 9))
    add_p(c, "Số cá\nnhân\nđược\nthăm hỏi,\nđộng\nviên\ntrong dịp\ntổ chức\nNgày\nhội/\nSố tiền\nthực hiện", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    c = table.cell(1, 10).merge(table.cell(2, 10))
    add_p(c, "Ước\nlượng\ntổng số\nngười\ntham\ndự\nNgày\nhội\ntrên địa\nbàn", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    cell_ld = table.cell(1, 11)
    for c_idx in range(12, 16):
        cell_ld.merge(table.cell(1, c_idx))
    add_p(cell_ld, "Tổng số lượt lãnh đạo các cấp\ntrong tỉnh, phường tham dự Ngày hội", align=WD_ALIGN_PARAGRAPH.CENTER, size=8, space_after=0, line_spacing=1.0)
    
    add_p(table.cell(2, 11), "Tổng\nsố", align=WD_ALIGN_PARAGRAPH.CENTER, size=7, space_after=0, line_spacing=1.0)
    add_p(table.cell(2, 12), "Lãnh\nđạo\ncấp\ntrung\nương\ntham\ndự\nngày\nhội\n(lượt\ndự)", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=7, space_after=0, line_spacing=1.0)
    add_p(table.cell(2, 13), "Lãnh\nđạo\nBTV,\nBCH\nCấp\ntỉnh\n(tổng\nsố lãnh\nđạo/\nlượt\ndự)", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=7, space_after=0, line_spacing=1.0)
    add_p(table.cell(2, 14), "Lãnh\nđạo\ncác cơ\nquan\ncấp sở\n(tổng\nsố lãnh\nđạo/\nlượt\ndự)", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=7, space_after=0, line_spacing=1.0)
    add_p(table.cell(2, 15), "Lãnh\nđạo\ncấp xã\ndự\nNgày\nhội\n(tổng\nsố lãnh\nđạo/\nlượt\ndự)", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=7, space_after=0, line_spacing=1.0)
    
    c = table.cell(1, 16).merge(table.cell(2, 16))
    add_p(c, "Tổng\nsố\ncông\ntrình\ndân\nsinh\nđược\nxây\ndựng\ntrong\ndịp tổ\nchức\nNgày\nhội", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    c = table.cell(1, 17).merge(table.cell(2, 17))
    add_p(c, "Tổng\nsố\ncác\nmô\nhình\nđược\nxây\ndựng\nvà\nnâng\ncao", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    c = table.cell(1, 18).merge(table.cell(2, 18))
    add_p(c, "Tổng\nsố nhà\nĐại\nđoàn\nkết\nđược\nxây\nmới\n(căn)/\ntổng\nkinh\nphí", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    c = table.cell(1, 19).merge(table.cell(2, 19))
    add_p(c, "Tổng số\nnhà Đại\nđoàn kết\nđược\nsửa\nchữa\n(căn)/\ntổng\nkinh phí", align=WD_ALIGN_PARAGRAPH.CENTER, size=7.5, space_after=0, line_spacing=1.0)
    
    # ROW 3: Numbering row (1..16)
    nums = [
        (0, 0, "1"),
        (1, 1, "2"),
        (2, 2, "3"),
        (3, 3, "4"),
        (4, 4, "5"),
        (5, 5, "6"),
        (6, 6, "7"),
        (7, 7, "8"),
        (8, 8, "9"),
        (9, 9, "10"),
        (10, 10, "11"),
        (11, 15, "12"),
        (16, 16, "13"),
        (17, 17, "14"),
        (18, 18, "15"),
        (19, 19, "16")
    ]
    for c_start, c_end, val in nums:
        cell_num = table.cell(3, c_start)
        if c_end > c_start:
            for ci in range(c_start + 1, c_end + 1):
                cell_num.merge(table.cell(3, ci))
        add_p(cell_num, val, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=8, space_after=0, line_spacing=1.0)
        
    # ROWS 4 & 5: Empty data rows
    for r_idx in [4, 5]:
        row = table.rows[r_idx]
        trPr = row._tr.get_or_add_trPr()
        trHeight = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="320" w:hRule="atLeast"/>')
        trPr.append(trHeight)
    
    # CRITICAL: Prevent Word from creating an empty landscape page after the table
    # Word requires an end paragraph; we make it 1pt size with zero spacing.
    end_p = doc.add_paragraph()
    end_p.paragraph_format.space_before = Pt(0)
    end_p.paragraph_format.space_after = Pt(0)
    end_p.paragraph_format.line_spacing = Pt(1)
    end_run = end_p.add_run("")
    end_run.font.size = Pt(1)
    
    out_path = "/Users/haisonnguyenphuc/Desktop/project/word/Ke_hoach_Ngay_hoi_Dai_doan_ket_2026.docx"
    doc.save(out_path)
    print("Document successfully created at:", out_path)

if __name__ == '__main__':
    generate_document()
